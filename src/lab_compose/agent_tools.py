"""Workspace and authoring primitives exposed through the MCP server."""

from __future__ import annotations

import subprocess
from pathlib import Path
from typing import Any

import yaml
from docx import Document

from console_gen.schemas import load_schema as load_screenshot_schema
from console_gen.schemas import validate_instance as validate_screenshot_instance
from otchet_compose.schemas import load_schema as load_report_schema
from otchet_compose.schemas import validate_instance as validate_report_instance

from .schemas import load_schema as load_pipeline_schema
from .schemas import validate_instance as validate_pipeline_instance


_SCHEMA_LOADERS = {
    "pipeline": load_pipeline_schema,
    "report": load_report_schema,
    "screenshots": load_screenshot_schema,
}

_SCHEMA_VALIDATORS = {
    "pipeline": validate_pipeline_instance,
    "report": validate_report_instance,
    "screenshots": validate_screenshot_instance,
}

_EXAMPLES: dict[str, dict[str, Any]] = {
    "pipeline": {
        "version": 1,
        "screenshots": {"project": "screenshot-compose.yml"},
        "report": {"config": "otchet-compose.yml"},
        "manifest": "build/lab-compose.manifest.json",
    },
    "screenshots": {
        "version": 1,
        "renders": {
            "terminal": {
                "input": "artifacts/terminal.log",
                "output": "build/terminal.png",
                "title": "Laboratory result",
                "frame": "windows",
            }
        },
    },
    "report": {
        "version": 1,
        "document": {"output": "build/report.docx", "toc": False},
        "content": [
            {"type": "heading", "text": "ХОД РАБОТЫ", "level": 1, "structural": True},
            {"type": "paragraph", "text": "Описание выполненной работы."},
            {"type": "figure", "caption": "Результат выполнения", "path": "build/terminal.png"},
        ],
    },
}


def get_schema(kind: str, version: int = 1) -> dict:
    """Return a compose JSON Schema by kind."""
    loader = _SCHEMA_LOADERS.get(kind)
    if loader is None:
        raise ValueError(f"Unknown schema kind: {kind}. Expected one of: {', '.join(sorted(_SCHEMA_LOADERS))}")
    return loader(version)


def get_example(kind: str) -> dict:
    """Return a minimal authoring example by kind."""
    example = _EXAMPLES.get(kind)
    if example is None:
        raise ValueError(f"Unknown example kind: {kind}. Expected one of: {', '.join(sorted(_EXAMPLES))}")
    return example


def list_workspace_files(workspace_root: str, pattern: str = "**/*", limit: int = 500) -> dict:
    """List files inside a workspace using a glob pattern."""
    root = _workspace_root(workspace_root)
    if not 1 <= limit <= 5000:
        raise ValueError("limit must be between 1 and 5000")
    paths = []
    for path in root.glob(pattern):
        resolved = path.resolve()
        if resolved.is_file() and resolved.is_relative_to(root):
            paths.append(resolved.relative_to(root).as_posix())
            if len(paths) >= limit:
                break
    return {"workspace_root": str(root), "files": paths, "truncated": len(paths) >= limit}


def read_text_file(workspace_root: str, path: str, max_bytes: int = 200_000) -> dict:
    """Read a UTF-8 text file inside a workspace."""
    if not 1 <= max_bytes <= 2_000_000:
        raise ValueError("max_bytes must be between 1 and 2000000")
    target = _workspace_path(workspace_root, path)
    size = target.stat().st_size
    if size > max_bytes:
        raise ValueError(f"File is too large: {size} bytes; max_bytes is {max_bytes}")
    return {"path": str(target), "content": target.read_text(encoding="utf-8-sig"), "size": size}


def read_document(workspace_root: str, path: str, max_chars: int = 200_000) -> dict:
    """Extract text from a UTF-8 text, DOCX, or PDF assignment inside a workspace."""
    if not 1 <= max_chars <= 2_000_000:
        raise ValueError("max_chars must be between 1 and 2000000")
    target = _workspace_path(workspace_root, path)
    suffix = target.suffix.lower()
    if suffix == ".docx":
        document = Document(target)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        content = "\n".join(parts)
    elif suffix == ".pdf":
        from pypdf import PdfReader

        content = "\n\n".join(page.extract_text() or "" for page in PdfReader(target).pages)
    else:
        content = target.read_text(encoding="utf-8-sig")
    truncated = len(content) > max_chars
    return {
        "path": str(target),
        "content": content[:max_chars],
        "characters": len(content),
        "truncated": truncated,
    }


def write_text_file(workspace_root: str, path: str, content: str, overwrite: bool = False) -> dict:
    """Write a UTF-8 text file inside a workspace."""
    target = _workspace_path(workspace_root, path, must_exist=False)
    _check_overwrite(target, overwrite)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8")
    return {"path": str(target), "size": target.stat().st_size}


def write_compose_config(
    workspace_root: str,
    path: str,
    kind: str,
    config: dict,
    overwrite: bool = False,
) -> dict:
    """Validate and write a pipeline, report, or screenshots YAML config."""
    validator = _SCHEMA_VALIDATORS.get(kind)
    if validator is None:
        raise ValueError(f"Unknown config kind: {kind}. Expected one of: {', '.join(sorted(_SCHEMA_VALIDATORS))}")
    version = config.get("version", 1)
    validator(config, version)
    content = yaml.safe_dump(config, allow_unicode=True, sort_keys=False)
    result = write_text_file(workspace_root, path, content, overwrite)
    return {**result, "kind": kind, "version": version}


def run_workspace_command(
    workspace_root: str,
    command: list[str],
    cwd: str = ".",
    timeout_seconds: int = 120,
    log_path: str | None = None,
    overwrite_log: bool = False,
    max_output_chars: int = 100_000,
) -> dict:
    """Run an argv-style command from a workspace directory and capture its output.

    The command is executed directly without shell parsing. The working directory
    is restricted to the workspace, but the child process otherwise has the same
    host permissions as the MCP server.
    """
    if not command or not all(isinstance(part, str) and part for part in command):
        raise ValueError("command must be a non-empty list of non-empty strings")
    if not 1 <= timeout_seconds <= 900:
        raise ValueError("timeout_seconds must be between 1 and 900")
    if not 1 <= max_output_chars <= 1_000_000:
        raise ValueError("max_output_chars must be between 1 and 1000000")
    working_directory = _workspace_path(workspace_root, cwd)
    if not working_directory.is_dir():
        raise ValueError(f"cwd is not a directory: {working_directory}")

    try:
        completed = subprocess.run(
            command,
            cwd=working_directory,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout_seconds,
            shell=False,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise ValueError(f"Command timed out after {timeout_seconds} seconds") from exc

    result = {
        "command": command,
        "cwd": str(working_directory),
        "exit_code": completed.returncode,
        "stdout": completed.stdout[:max_output_chars],
        "stderr": completed.stderr[:max_output_chars],
        "stdout_truncated": len(completed.stdout) > max_output_chars,
        "stderr_truncated": len(completed.stderr) > max_output_chars,
    }
    if log_path:
        log_content = completed.stdout
        if completed.stderr:
            log_content += ("\n" if log_content and not log_content.endswith("\n") else "") + completed.stderr
        written = write_text_file(workspace_root, log_path, log_content, overwrite_log)
        result["log_path"] = written["path"]
    return result


def _workspace_root(workspace_root: str) -> Path:
    root = Path(workspace_root).resolve()
    if not root.is_dir():
        raise ValueError(f"workspace_root is not a directory: {root}")
    return root


def _workspace_path(workspace_root: str, path: str, *, must_exist: bool = True) -> Path:
    root = _workspace_root(workspace_root)
    candidate = Path(path)
    target = candidate.resolve() if candidate.is_absolute() else (root / candidate).resolve()
    if not target.is_relative_to(root):
        raise ValueError(f"Path is outside workspace_root: {target}")
    if must_exist and not target.exists():
        raise FileNotFoundError(f"Path does not exist: {target}")
    return target


def _check_overwrite(path: Path, overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"File already exists; set overwrite=true to replace: {path}")
