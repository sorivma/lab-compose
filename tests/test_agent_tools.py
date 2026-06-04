import sys
from pathlib import Path

import pytest
import yaml
from docx import Document

from lab_compose.agent_tools import (
    get_example,
    get_schema,
    list_workspace_files,
    read_document,
    read_text_file,
    run_workspace_command,
    write_compose_config,
    write_text_file,
)


def test_schemas_and_examples_cover_all_config_kinds():
    for kind in ("pipeline", "report", "screenshots"):
        assert get_schema(kind)["type"] == "object"
        assert get_example(kind)["version"] == 1


def test_workspace_text_tools_reject_escape(tmp_path):
    written = write_text_file(str(tmp_path), "src/app.txt", "hello")
    assert Path(written["path"]).read_text(encoding="utf-8") == "hello"
    assert read_text_file(str(tmp_path), "src/app.txt")["content"] == "hello"
    assert list_workspace_files(str(tmp_path), "**/*.txt")["files"] == ["src/app.txt"]

    with pytest.raises(ValueError, match="outside workspace_root"):
        write_text_file(str(tmp_path), "../outside.txt", "no")


def test_read_document_extracts_docx_assignment(tmp_path):
    assignment = Document()
    assignment.add_heading("Laboratory 1", level=1)
    assignment.add_paragraph("Implement and verify the solution.")
    assignment.save(tmp_path / "assignment.docx")

    result = read_document(str(tmp_path), "assignment.docx")
    assert "Laboratory 1" in result["content"]
    assert "Implement and verify" in result["content"]
    assert result["truncated"] is False


def test_write_compose_config_validates_before_writing(tmp_path):
    result = write_compose_config(
        str(tmp_path),
        "screenshot-compose.yml",
        "screenshots",
        get_example("screenshots"),
    )
    config = yaml.safe_load(Path(result["path"]).read_text(encoding="utf-8"))
    assert "terminal" in config["renders"]

    with pytest.raises(ValueError, match="JSON Schema"):
        write_compose_config(str(tmp_path), "invalid.yml", "report", {"version": 1})
    assert not (tmp_path / "invalid.yml").exists()


def test_run_workspace_command_captures_and_saves_output(tmp_path):
    result = run_workspace_command(
        str(tmp_path),
        [sys.executable, "-c", "print('laboratory complete')"],
        log_path="artifacts/run.log",
    )
    assert result["exit_code"] == 0
    assert result["stdout"] == "laboratory complete\n"
    assert result["stdout_truncated"] is False
    assert (tmp_path / "artifacts" / "run.log").read_text(encoding="utf-8") == "laboratory complete\n"


def test_run_workspace_command_restricts_cwd(tmp_path):
    with pytest.raises(ValueError, match="outside workspace_root"):
        run_workspace_command(str(tmp_path), [sys.executable, "--version"], cwd="..")
